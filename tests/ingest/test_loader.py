# tests/ingest/test_loader.py
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from docx import Document as DocxDocument

from ingest.loader import load_text


def test_load_docx_joins_paragraphs(tmp_path):
    path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("Chính sách nghỉ phép.")
    doc.add_paragraph("Nhân viên được nghỉ 12 ngày mỗi năm.")
    doc.save(path)

    text = load_text(path)

    assert "Chính sách nghỉ phép." in text
    assert "Nhân viên được nghỉ 12 ngày mỗi năm." in text


def test_load_pdf_joins_page_text(tmp_path):
    path = tmp_path / "sample.pdf"
    path.write_bytes(b"%PDF-1.4 fake")  # content irrelevant, PdfReader is mocked

    fake_page_1 = MagicMock()
    fake_page_1.extract_text.return_value = "Trang 1"
    fake_page_2 = MagicMock()
    fake_page_2.extract_text.return_value = "Trang 2"

    with patch("ingest.loader.PdfReader") as mock_reader:
        mock_reader.return_value.pages = [fake_page_1, fake_page_2]
        text = load_text(path)

    assert text == "Trang 1\nTrang 2"


def test_rejects_unsupported_extension(tmp_path):
    path = tmp_path / "sample.txt"
    path.write_text("hello")

    with pytest.raises(ValueError, match="Unsupported file type"):
        load_text(path)
